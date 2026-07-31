"""Excel 数据源实现（首版）。

使用 openpyxl 读取 backend/data/cicada_points.xlsx，首次访问时缓存到内存。
"""
from __future__ import annotations

from typing import List, Optional

from openpyxl import load_workbook

from app.config import EXCEL_FILE
from app.data.repository import CicadaRepository
from app.schemas import CicadaPoint, CicadaStats, SpeciesInfo

# 北京四种常见鸣蝉（与 Excel 的 species 字段呼应，用于物种面板）
_SPECIES = [
    SpeciesInfo(name="黑蚱蝉", alias="麻季鸟儿", habitat="城区最常见，体型最大全黑", season="6–8月"),
    SpeciesInfo(name="蒙古寒蝉", alias="伏天儿", habitat="山区偏多，偏绿", season="7–9月"),
    SpeciesInfo(name="鸣鸣蝉", alias="斑透翅蝉", habitat="山区，蓝绿纹", season="8月集中"),
    SpeciesInfo(name="蟪蛄", alias="小热热儿", habitat="最小，树皮伪装", season="5月最早"),
]


class ExcelCicadaRepository(CicadaRepository):
    def __init__(self, path: str = EXCEL_FILE):
        self._path = path
        self._cache: Optional[List[CicadaPoint]] = None

    def _load(self) -> List[CicadaPoint]:
        if self._cache is not None:
            return self._cache
        wb = load_workbook(self._path, read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
        if not rows:
            self._cache = []
            return self._cache
        header = [str(h).strip() if h is not None else "" for h in rows[0]]
        points: List[CicadaPoint] = []
        for raw in rows[1:]:
            if raw is None or all(c is None for c in raw):
                continue
            row = {header[i]: raw[i] for i in range(min(len(header), len(raw)))}
            # 跳过空行 / 缺关键字段
            if row.get("id") in (None, "") or row.get("lng") in (None, ""):
                continue
            points.append(CicadaPoint.from_row(row))
        self._cache = points
        return self._cache

    def list_points(self) -> List[CicadaPoint]:
        return self._load()

    def get_point(self, point_id: int) -> Optional[CicadaPoint]:
        for p in self._load():
            if p.id == point_id:
                return p
        return None

    def list_species(self) -> List[SpeciesInfo]:
        return _SPECIES

    def get_stats(self) -> CicadaStats:
        points = self._load()
        levels = {"高": 0, "中": 0, "低": 0}
        by_district: dict[str, int] = {}
        for p in points:
            levels[p.abundance] = levels.get(p.abundance, 0) + 1
            by_district[p.district] = by_district.get(p.district, 0) + 1
        return CicadaStats(
            total=len(points),
            districts=len(by_district),
            levels=levels,
            by_district=by_district,
        )
