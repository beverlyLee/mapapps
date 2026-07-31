"""Excel 数据维护工具（CLI）。

更新 backend/data/cicada_points.xlsx 的命令行工具，配合后端 mtime 热加载机制，
增删改后无需重启服务，刷新浏览器页面即可看到新数据。

用法:
    # 列出全部点位
    python update_excel.py list

    # 添加点位（id 自动取 max+1；abundance 默认"中"；is_famous 默认随 abundance）
    python update_excel.py add \
        --name "示例公园" \
        --province "北京市" --city "北京市" --district "朝阳区" \
        --lng 116.396 --lat 39.992 \
        --species "黑蚱蝉·蟪蛄" \
        --peak-season "6–8月" \
        --abundance 高

    # 修改点位（只传需要改的字段）
    python update_excel.py update 22 --abundance 高 --is-famous TRUE

    # 删除点位
    python update_excel.py delete 22

    # 重置为种子数据（gen_excel.py 中的 45 条）
    python update_excel.py reset

    # 导出为 CSV（便于在表格软件中查看）
    python update_excel.py export

字段说明:
    id           整数，唯一
    province     省份/直辖市
    city         城市
    district     区/县
    town         乡镇/村级（可空）
    name         点位名称
    lng, lat     GCJ-02 坐标（高德坐标系）
    species      蝉种，多个用 · 分隔
    peak_season  旺季（如 6–8月）
    abundance    高 / 中 / 低
    is_famous    TRUE / FALSE（留空时自动 = (abundance == "高")）
"""
from __future__ import annotations

import argparse
import csv
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook

DATA_DIR = Path(__file__).resolve().parent / "data"
EXCEL_FILE = DATA_DIR / "cicada_points.xlsx"
HEADER = ["id", "province", "city", "district", "town",
          "name", "lng", "lat", "species", "peak_season", "abundance", "is_famous"]

from gen_excel import ROWS as SEED_ROWS  # noqa: E402


def load_rows() -> list[dict]:
    if not EXCEL_FILE.exists():
        return []
    wb = load_workbook(EXCEL_FILE, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return []
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    out = []
    for raw in rows[1:]:
        if raw is None or all(c is None for c in raw):
            continue
        row = {header[i]: raw[i] for i in range(min(len(header), len(raw)))}
        if row.get("id") in (None, ""):
            continue
        out.append(_normalize(row))
    return out


def save_rows(rows: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "cicada_points"
    ws.append(HEADER)
    for r in sorted(rows, key=lambda x: int(x["id"])):
        ws.append([r[k] for k in HEADER])
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(EXCEL_FILE)


def _normalize(row: dict) -> dict:
    abundance = str(row.get("abundance") or "中").strip() or "中"
    is_famous = row.get("is_famous")
    if is_famous is None or is_famous == "":
        is_famous = abundance == "高"
    else:
        is_famous = bool(is_famous) if isinstance(is_famous, bool) \
            else str(is_famous).strip().upper() in ("TRUE", "1", "YES")
    return {
        "id": int(row["id"]),
        "province": str(row.get("province") or "北京市"),
        "city": str(row.get("city") or row.get("province") or "北京市"),
        "district": str(row.get("district") or ""),
        "town": str(row.get("town") or ""),
        "name": str(row.get("name") or ""),
        "lng": float(row.get("lng") or 0),
        "lat": float(row.get("lat") or 0),
        "species": str(row.get("species") or ""),
        "peak_season": str(row.get("peak_season") or ""),
        "abundance": abundance,
        "is_famous": is_famous,
    }


def cmd_list(_args) -> int:
    rows = load_rows()
    if not rows:
        print(f"（空）Excel 文件不存在或无数据：{EXCEL_FILE}")
        return 0
    print(f"共 {len(rows)} 条点位  ←  {EXCEL_FILE}")
    print("-" * 120)
    print(f"{'id':>3}  {'名称':<24} {'省':<6} {'市':<6} {'区':<8} {'乡镇':<10} {'lng':>10} {'lat':>10} {'丰富度':<4} {'名':<3}  蝉种")
    print("-" * 120)
    for r in rows:
        town = (r["town"] or "")[:10]
        print(f"{r['id']:>3}  {r['name'][:24]:<24} {r['province']:<6} {r['city']:<6} "
              f"{r['district']:<8} {town:<10} {r['lng']:>10.4f} {r['lat']:>10.4f}  "
              f"{r['abundance']:<4} {'是' if r['is_famous'] else '否':<3}  {r['species']}")
    return 0


def _next_id(rows: list[dict]) -> int:
    return (max((int(r["id"]) for r in rows), default=0) + 1)


def cmd_add(args) -> int:
    rows = load_rows()
    new_id = args.id if args.id is not None else _next_id(rows)
    if any(int(r["id"]) == new_id for r in rows):
        print(f"错误：id={new_id} 已存在", file=sys.stderr)
        return 1
    abundance = args.abundance
    if abundance not in ("高", "中", "低"):
        print(f"错误：abundance 必须是 高/中/低", file=sys.stderr)
        return 1
    is_famous = args.is_famous
    if is_famous is None:
        is_famous = abundance == "高"
    row = {
        "id": new_id,
        "province": args.province or "北京市",
        "city": args.city or args.province or "北京市",
        "district": args.district or "",
        "town": args.town or "",
        "name": args.name,
        "lng": args.lng,
        "lat": args.lat,
        "species": args.species or "",
        "peak_season": args.peak_season or "",
        "abundance": abundance,
        "is_famous": is_famous,
    }
    rows.append(_normalize(row))
    save_rows(rows)
    print(f"已添加：id={new_id}  {args.name}  →  {EXCEL_FILE}")
    return 0


def cmd_update(args) -> int:
    rows = load_rows()
    target = next((r for r in rows if int(r["id"]) == args.id), None)
    if target is None:
        print(f"错误：未找到 id={args.id}", file=sys.stderr)
        return 1
    updates = {k: v for k, v in vars(args).items()
               if k != "id" and v is not None}
    if "abundance" in updates and updates["abundance"] not in ("高", "中", "低"):
        print(f"错误：abundance 必须是 高/中/低", file=sys.stderr)
        return 1
    target.update(updates)
    target = _normalize(target)
    rows = [target if int(r["id"]) == args.id else r for r in rows]
    save_rows(rows)
    print(f"已更新：id={args.id}  →  {target}")
    return 0


def cmd_delete(args) -> int:
    rows = load_rows()
    before = len(rows)
    rows = [r for r in rows if int(r["id"]) != args.id]
    if len(rows) == before:
        print(f"错误：未找到 id={args.id}", file=sys.stderr)
        return 1
    save_rows(rows)
    print(f"已删除：id={args.id}（剩 {len(rows)} 条）→  {EXCEL_FILE}")
    return 0


def cmd_reset(_args) -> int:
    rows = []
    for r in SEED_ROWS:
        rid, prov, city, dist, town, name, lng, lat, sp, peak, ab = r
        rows.append({
            "id": rid, "province": prov, "city": city,
            "district": dist, "town": town,
            "name": name, "lng": lng, "lat": lat,
            "species": sp, "peak_season": peak,
            "abundance": ab, "is_famous": ab == "高",
        })
    save_rows(rows)
    print(f"已重置为种子数据（{len(rows)} 条）→  {EXCEL_FILE}")
    return 0


def cmd_export(_args) -> int:
    rows = load_rows()
    csv_path = DATA_DIR / "cicada_points.csv"
    with open(csv_path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(HEADER)
        for r in rows:
            w.writerow([r[k] for k in HEADER])
    print(f"已导出 CSV：{csv_path}（{len(rows)} 行）")
    return 0


def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(
        description=" cicada_points.xlsx 维护工具（改完刷新浏览器即生效）",
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    sub = p.add_subparsers(dest="cmd", required=True)

    sub.add_parser("list", help="列出全部点位").set_defaults(func=cmd_list)
    sub.add_parser("reset", help="重置为种子数据").set_defaults(func=cmd_reset)
    sub.add_parser("export", help="导出为 CSV").set_defaults(func=cmd_export)

    pa = sub.add_parser("add", help="添加点位")
    pa.add_argument("--id", type=int, default=None)
    pa.add_argument("--province", default="北京市")
    pa.add_argument("--city", default=None, help="默认同 province")
    pa.add_argument("--district", default="")
    pa.add_argument("--town", default="")
    pa.add_argument("--name", required=True)
    pa.add_argument("--lng", type=float, required=True)
    pa.add_argument("--lat", type=float, required=True)
    pa.add_argument("--species", default="")
    pa.add_argument("--peak-season", default="", dest="peak_season")
    pa.add_argument("--abundance", default="中", choices=["高", "中", "低"])
    pa.add_argument("--is-famous", default=None, dest="is_famous",
                    type=lambda s: s.upper() == "TRUE")
    pa.set_defaults(func=cmd_add)

    pu = sub.add_parser("update", help="修改点位")
    pu.add_argument("id", type=int)
    pu.add_argument("--province")
    pu.add_argument("--city")
    pu.add_argument("--district")
    pu.add_argument("--town")
    pu.add_argument("--name")
    pu.add_argument("--lng", type=float)
    pu.add_argument("--lat", type=float)
    pu.add_argument("--species")
    pu.add_argument("--peak-season", dest="peak_season")
    pu.add_argument("--abundance", choices=["高", "中", "低"])
    pu.add_argument("--is-famous", dest="is_famous",
                    type=lambda s: s.upper() == "TRUE")
    pu.set_defaults(func=cmd_update)

    pd = sub.add_parser("delete", help="删除点位")
    pd.add_argument("id", type=int)
    pd.set_defaults(func=cmd_delete)
    return p


def main() -> int:
    args = build_parser().parse_args()
    return args.func(args)


if __name__ == "__main__":
    sys.exit(main())
